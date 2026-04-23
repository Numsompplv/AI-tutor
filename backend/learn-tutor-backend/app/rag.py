import os
import faiss
import numpy as np
import pickle
from typing import List, Dict, Any
from pathlib import Path

# Document parsers
import pypdf
from docx import Document as DocxDocument
from pptx import Presentation

# Embeddings
from sentence_transformers import SentenceTransformer

# Ollama for LLM
import requests
import json

# ── Config ──
FAISS_DIR = "faiss_indexes"
os.makedirs(FAISS_DIR, exist_ok=True)

CHUNK_SIZE = 500
CHUNK_OVERLAP = 50
TOP_K = 3

OLLAMA_URL = "http://localhost:11434/api/generate"
OLLAMA_MODEL = "llama3.2"

# ── Load embedding model once ──
print("Loading embedding model...")
embedder = SentenceTransformer("all-MiniLM-L6-v2")
print("✅ Embedding model ready")


# ══════════════════════════════════════════
# TEXT EXTRACTION
# ══════════════════════════════════════════

def extract_text_from_pdf(file_path: str) -> List[Dict]:
    """Extract text from PDF with page numbers."""
    chunks_meta = []
    reader = pypdf.PdfReader(file_path)
    for page_num, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if text.strip():
            chunks_meta.append({"text": text, "page": page_num})
    return chunks_meta


def extract_text_from_docx(file_path: str) -> List[Dict]:
    """Extract text from DOCX."""
    doc = DocxDocument(file_path)
    full_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    return [{"text": full_text, "page": 1}]


def extract_text_from_pptx(file_path: str) -> List[Dict]:
    """Extract text from PPTX with slide numbers."""
    prs = Presentation(file_path)
    chunks_meta = []
    for slide_num, slide in enumerate(prs.slides, 1):
        text = ""
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text += shape.text + "\n"
        if text.strip():
            chunks_meta.append({"text": text, "page": slide_num})
    return chunks_meta


def extract_text_from_txt(file_path: str) -> List[Dict]:
    """Extract text from TXT."""
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read()
    return [{"text": text, "page": 1}]


def extract_text(file_path: str, file_type: str) -> List[Dict]:
    """Route to correct extractor based on file type."""
    extractors = {
        "pdf": extract_text_from_pdf,
        "docx": extract_text_from_docx,
        "pptx": extract_text_from_pptx,
        "txt": extract_text_from_txt,
    }
    extractor = extractors.get(file_type.lower())
    if not extractor:
        raise ValueError(f"Unsupported file type: {file_type}")
    return extractor(file_path)


# ══════════════════════════════════════════
# TEXT CHUNKING
# ══════════════════════════════════════════

def chunk_text(text: str, page: int, doc_name: str) -> List[Dict]:
    """Split text into overlapping chunks."""
    chunks = []
    start = 0
    while start < len(text):
        end = start + CHUNK_SIZE
        chunk = text[start:end]
        if chunk.strip():
            chunks.append({
                "text": chunk.strip(),
                "page": page,
                "doc": doc_name,
            })
        start += CHUNK_SIZE - CHUNK_OVERLAP
    return chunks


# ══════════════════════════════════════════
# FAISS INDEX MANAGEMENT
# ══════════════════════════════════════════

def get_index_path(notebook_id: str):
    return os.path.join(FAISS_DIR, f"{notebook_id}.index")


def get_chunks_path(notebook_id: str):
    return os.path.join(FAISS_DIR, f"{notebook_id}.pkl")


def load_index(notebook_id: str):
    """Load FAISS index and chunks for a notebook."""
    index_path = get_index_path(notebook_id)
    chunks_path = get_chunks_path(notebook_id)
    if not os.path.exists(index_path):
        return None, []
    index = faiss.read_index(index_path)
    with open(chunks_path, "rb") as f:
        chunks = pickle.load(f)
    return index, chunks


def save_index(notebook_id: str, index, chunks: List[Dict]):
    """Save FAISS index and chunks for a notebook."""
    faiss.write_index(index, get_index_path(notebook_id))
    with open(get_chunks_path(notebook_id), "wb") as f:
        pickle.dump(chunks, f)


# ══════════════════════════════════════════
# DOCUMENT INGESTION
# ══════════════════════════════════════════

def ingest_document(notebook_id: str, file_path: str, file_type: str, doc_name: str) -> int:
    """
    Process a document and add it to the notebook's FAISS index.
    Returns the number of chunks created.
    """
    # Extract text
    pages = extract_text(file_path, file_type)

    # Chunk text
    all_chunks = []
    for page_data in pages:
        chunks = chunk_text(page_data["text"], page_data["page"], doc_name)
        all_chunks.extend(chunks)

    if not all_chunks:
        return 0

    # Generate embeddings
    texts = [c["text"] for c in all_chunks]
    embeddings = embedder.encode(texts, show_progress_bar=False)
    embeddings = np.array(embeddings, dtype=np.float32)

    # Load or create FAISS index
    index, existing_chunks = load_index(notebook_id)
    if index is None:
        dim = embeddings.shape[1]
        index = faiss.IndexFlatL2(dim)

    # Add to index
    index.add(embeddings)
    existing_chunks.extend(all_chunks)

    # Save updated index
    save_index(notebook_id, index, existing_chunks)

    return len(all_chunks)


# ══════════════════════════════════════════
# RETRIEVAL
# ══════════════════════════════════════════

def retrieve_chunks(notebook_id: str, query: str, top_k: int = TOP_K) -> List[Dict]:
    """Find the most relevant chunks for a query."""
    index, chunks = load_index(notebook_id)
    if index is None or len(chunks) == 0:
        return []

    query_embedding = embedder.encode([query], show_progress_bar=False)
    query_embedding = np.array(query_embedding, dtype=np.float32)

    distances, indices = index.search(query_embedding, min(top_k, len(chunks)))

    results = []
    for i, idx in enumerate(indices[0]):
        if idx != -1:
            chunk = chunks[idx].copy()
            chunk["score"] = float(distances[0][i])
            results.append(chunk)

    return results


# ══════════════════════════════════════════
# LLM GENERATION
# ══════════════════════════════════════════

def ask_ollama(prompt: str) -> str:
    """Send a prompt to Ollama and get a response."""
    try:
        response = requests.post(
            OLLAMA_URL,
            json={
                "model": OLLAMA_MODEL,
                "prompt": prompt,
                "stream": False,
            },
            timeout=60,
        )
        response.raise_for_status()
        return response.json().get("response", "").strip()
    except requests.exceptions.ConnectionError:
        return "❌ Ollama is not running. Please start it with: ollama serve"
    except Exception as e:
        return f"❌ Error: {str(e)}"


# ══════════════════════════════════════════
# RAG — MAIN FUNCTION
# ══════════════════════════════════════════

def rag_answer(notebook_id: str, question: str, conversation_history: 
               List[Dict] = None) -> Dict:
    """
    Full RAG pipeline:
    1. Retrieve relevant chunks
    2. Build prompt with context
    3. Generate answer with Ollama
    4. Return answer + citations
    """
    chunks = retrieve_chunks(notebook_id, question)

    if not chunks:
        prompt = f"""You are a friendly and helpful AI tutor. 
        Answer the following question clearly and helpfully.

Question: {question}

Answer:"""
        answer = ask_ollama(prompt)
        return {
            "answer": answer,
            "citations": [],
            "mode": "general",
        }

    context = "\n\n".join([
        f"[Source: {c['doc']}, Page {c['page']}]\n{c['text']}"
        for c in chunks
    ])

    prompt = f"""You are a friendly and helpful AI tutor. Answer the student's question based ONLY on the provided course materials below. Be clear, educational, and encouraging. If the answer is not in the materials, say so honestly.

Course Materials:
{context}

Student Question: {question}

Answer (be clear and educational):"""

    answer = ask_ollama(prompt)

    # Build citations
    citations = [
        {
            "doc": c["doc"],
            "page": c["page"],
            "text": c["text"][:150] + "..." if len(c["text"]) > 150 else c["text"],
        }
        for c in chunks
    ]

    return {
        "answer": answer,
        "citations": citations,
        "mode": "rag",
    }


def free_learning_answer(topic: str, question: str) -> Dict:
    """Answer using general knowledge (no documents)."""
    prompt = f"""You are a friendly and helpful AI tutor teaching a student about: {topic}

Student Question: {question}

Give a clear, educational answer. Use examples where helpful. Be encouraging and easy to understand.

Answer:"""

    answer = ask_ollama(prompt)
    return {
        "answer": answer,
        "citations": [],
        "mode": "general",
    }